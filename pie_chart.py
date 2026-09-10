from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference

# Open the workbook
wb = load_workbook("students.xlsx")

# Select Students sheet
sheet = wb["Students"]

# Create a pie chart
chart = PieChart()

# Find the Marks column
marks_column = None

for cell in sheet[1]:
    if cell.value == "Marks":
        marks_column = cell.column
        break

if marks_column is None:
    print("Marks column not found")

else:
    # Select Marks data
    data = Reference(
        sheet,
        min_col=marks_column,
        min_row=1,
        max_row=sheet.max_row
    )

    # Select student names
    categories = Reference(
        sheet,
        min_col=1,
        min_row=2,
        max_row=sheet.max_row
    )

    # Add marks to chart
    chart.add_data(data, titles_from_data=True)

    # Add student names
    chart.set_categories(categories)

    # Chart title
    chart.title = "Student Marks - Pie Chart"

    # Add chart to worksheet
    sheet.add_chart(chart, "K38")

    # Save workbook
    wb.save("students.xlsx")

    print("Pie chart created successfully")