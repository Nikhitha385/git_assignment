from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import PatternFill

wb = load_workbook("students.xlsx")
sheet = wb["Students"]

# Green color for marks above 90
green_fill = PatternFill(
    start_color="00FF00",
    end_color="00FF00",
    fill_type="solid"
)

# Yellow color for marks below 80
yellow_fill = PatternFill(
    start_color="FFFF00",
    end_color="FFFF00",
    fill_type="solid"
)

# Above 90 = Green
green_rule = CellIsRule(
    operator="greaterThan",
    formula=["90"],
    fill=green_fill
)

# Below 80 = Yellow
yellow_rule = CellIsRule(
    operator="lessThan",
    formula=["80"],
    fill=yellow_fill
)

# Data bar = Blue
data_bar_rule = DataBarRule(
    start_type="min",
    end_type="max",
    color="638EC6",
    showValue=True
)

# Apply rules to Marks
sheet.conditional_formatting.add("B2:B15", green_rule)
sheet.conditional_formatting.add("B2:B15", yellow_rule)
sheet.conditional_formatting.add("B2:B15", data_bar_rule)

wb.save("students.xlsx")

print("Marks formatting completed successfully")