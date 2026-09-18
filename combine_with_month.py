from openpyxl import Workbook, load_workbook

files = {
    "January": "January.xlsx",
    "February": "February.xlsx",
    "March": "March.xlsx"
}

new_wb = Workbook()
new_sheet = new_wb.active

new_sheet.append(["Month", "Name", "Marks"])

for month, file in files.items():
    wb = load_workbook(file)
    sheet = wb.active

    for row in sheet.iter_rows(min_row=2, values_only=True):
        new_sheet.append([month, row[0], row[1]])

new_wb.save("All_Students_With_Month.xlsx")

print("Files combined with month successfully")