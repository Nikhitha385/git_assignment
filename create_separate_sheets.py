from openpyxl import Workbook, load_workbook

files = {
    "January": "January.xlsx",
    "February": "February.xlsx",
    "March": "March.xlsx"
}

new_wb = Workbook()

# Remove the default sheet
new_wb.remove(new_wb.active)

for month, file in files.items():

    wb = load_workbook(file)
    sheet = wb.active

    new_sheet = new_wb.create_sheet(month)

    for row in sheet.iter_rows(values_only=True):
        new_sheet.append(row)

new_wb.save("Monthly_Report.xlsx")

print("Separate sheets created successfully")