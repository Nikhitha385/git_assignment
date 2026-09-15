from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

wb = load_workbook("students.xlsx")
sheet = wb["Students"]

# Bar Chart
bar_chart = BarChart()

data = Reference(sheet, min_col=2, min_row=1, max_row=15)
names = Reference(sheet, min_col=1, min_row=2, max_row=15)

bar_chart.add_data(data, titles_from_data=True)
bar_chart.set_categories(names)

bar_chart.title = "Student Marks - Bar Chart"
bar_chart.y_axis.title = "Marks"
bar_chart.x_axis.title = "Students"

sheet.add_chart(bar_chart, "H2")


# Line Chart
line_chart = LineChart()

line_chart.add_data(data, titles_from_data=True)
line_chart.set_categories(names)

line_chart.title = "Student Marks - Line Chart"
line_chart.y_axis.title = "Marks"
line_chart.x_axis.title = "Students"

sheet.add_chart(line_chart, "H20")


# Pie Chart
pie_chart = PieChart()

pie_chart.add_data(data, titles_from_data=True)
pie_chart.set_categories(names)

pie_chart.title = "Student Marks - Pie Chart"

sheet.add_chart(pie_chart, "H38")


wb.save("students.xlsx")

print("Bar, Line and Pie charts restored successfully")