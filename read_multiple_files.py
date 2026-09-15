from openpyxl import load_workbook

files = ["January.xlsx", "February.xlsx", "March.xlsx"]

for file in files:
    wb = load_workbook(file)
    sheet = wb.active

    print("\n", file)

    for row in sheet.iter_rows(values_only=True):
        print(row)