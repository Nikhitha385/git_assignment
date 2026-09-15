from openpyxl import load_workbook
from openpyxl.styles import PatternFill

wb = load_workbook("students.xlsx")
sheet = wb["Students"]

# Remove normal fill from column D
for row in range(1, 17):
    sheet[f"D{row}"].fill = PatternFill(fill_type=None)

# Remove conditional formatting applied to column D
remove_rules = []

for cf_range in sheet.conditional_formatting._cf_rules:
    if "D" in str(cf_range):
        remove_rules.append(cf_range)

for cf_range in remove_rules:
    del sheet.conditional_formatting._cf_rules[cf_range]

wb.save("students.xlsx")

print("Yellow formatting removed from column D")