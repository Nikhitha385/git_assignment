from openpyxl import Workbook, load_workbook

files = ["January.xlsx", "February.xlsx", "March.xlsx"]

new_wb = Workbook()
new_sheet = new_wb.active

new_sheet.append(["Name", "Marks"])

for file in files:
    wb = load_workbook(file)
    sheet = wb.active

    for row in sheet.iter_rows(min_row=2, values_only=True):
        new_sheet.append(row)

new_wb.save("All_Students.xlsx")

print("All files combined successfully")