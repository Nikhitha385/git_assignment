from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

# Open the workbook
wb = load_workbook("students.xlsx")

# Select Students sheet
sheet = wb["Students"]

# Create a line chart
chart = LineChart()

# Find the Marks column
marks_column = None

for cell in sheet[1]:
    if cell.value == "Marks":
        marks_column = cell.column
        break

if marks_column is None:
    print("Marks column not found")

else:
    # Select Marks data
    data = Reference(
        sheet,
        min_col=marks_column,
        min_row=1,
        max_row=sheet.max_row
    )

    # Select student names
    categories = Reference(
        sheet,
        min_col=1,
        min_row=2,
        max_row=sheet.max_row
    )

    # Add data to chart
    chart.add_data(data, titles_from_data=True)

    # Add student names
    chart.set_categories(categories)

    # Chart title
    chart.title = "Student Marks - Line Chart"

    # Axis titles
    chart.y_axis.title = "Marks"
    chart.x_axis.title = "Students"

    # Place chart in worksheet
    sheet.add_chart(chart, "K20")

    # Save workbook
    wb.save("students.xlsx")

    print("Line chart created successfully")