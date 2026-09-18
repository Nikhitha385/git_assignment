from openpyxl import load_workbook

old_wb = load_workbook("students_old.xlsx")
new_wb = load_workbook("students_new.xlsx")

old_sheet = old_wb.active
new_sheet = new_wb.active

for row in range(2, old_sheet.max_row + 1):

    name = old_sheet[f"A{row}"].value
    old_marks = old_sheet[f"B{row}"].value
    new_marks = new_sheet[f"B{row}"].value

    if old_marks != new_marks:
        print(name, ":", old_marks, "->", new_marks)