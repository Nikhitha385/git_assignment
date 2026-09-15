from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

wb = load_workbook("students.xlsx")
sheet = wb["Students"]

red_fill = PatternFill(
    start_color="FF0000",
    end_color="FF0000",
    fill_type="solid"
)

rule = CellIsRule(
    operator="greaterThan",
    formula=["90"],
    fill=red_fill
)

sheet.conditional_formatting.add("D2:D16", rule)

wb.save("students.xlsx")

print("Conditional formatting added successfully")