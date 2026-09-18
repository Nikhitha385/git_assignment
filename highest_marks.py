from openpyxl import load_workbook
wb = load_workbook("All_Students_With_Month.xlsx")
sheet = wb.active
highest = 0
student = ""
for row in range(2, sheet.max_row + 1):
    name = sheet[f"B{row}"].value
    marks = sheet[f"C{row}"].value

    if marks > highest:
        highest = marks
        student = name

print("Highest Marks:", highest)
print("Student:", student)