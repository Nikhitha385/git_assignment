from openpyxl import Workbook, load_workbook

# Open existing Excel file
wb = load_workbook("students_new.xlsx")
sheet = wb.active

# Create a new workbook
report_wb = Workbook()
report_sheet = report_wb.active

# Add headers
report_sheet.append(["Name", "Marks", "Result"])

# Read student data
for row in range(2, sheet.max_row + 1):

    name = sheet[f"A{row}"].value
    marks = sheet[f"B{row}"].value

    if marks >= 40:
        result = "Pass"
    else:
        result = "Fail"

    report_sheet.append([name, marks, result])

# Save report
report_wb.save("Student_Report.xlsx")

print("Student report generated successfully")