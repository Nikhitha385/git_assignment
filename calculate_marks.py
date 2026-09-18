from openpyxl import load_workbook

wb = load_workbook("All_Students_With_Month.xlsx")
sheet = wb.active

total = 0
count = 0

for row in range(2, sheet.max_row + 1):
    marks = sheet[f"C{row}"].value
    total = total + marks
    count = count + 1

average = total / count

print("Total Marks:", total)
print("Average Marks:", average)